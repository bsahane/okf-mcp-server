"""Docling adapter tests; run by the ingestion CI job (`pytest -m docling`).

Skipped when the `ingest` extra is not installed. DOCX and XLSX conversion
needs no model downloads; PDF is exercised manually on real files.
"""

from types import SimpleNamespace

import pytest

pytest.importorskip("docling")

from okf_mcp_server.src.ingest.docling_adapter import extract_with_docling  # noqa: E402

pytestmark = pytest.mark.docling


@pytest.fixture(autouse=True)
def mock_imports():
    """Override conftest's sys.modules mocking: these tests run the real stack."""
    yield


def test_docx_sections_and_tables(tmp_path):
    from docx import Document

    path = tmp_path / "terms.docx"
    doc = Document()
    doc.add_heading("Warranty Terms", 0)
    doc.add_heading("Coverage", 1)
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "SKU", "Months"
    table.rows[1].cells[0].text, table.rows[1].cells[1].text = "SKU-4471", "36"
    doc.add_heading("Claims", 1)
    doc.add_paragraph("Raise claims within 14 days.")
    doc.save(path)

    result = extract_with_docling(path)
    titles = [s.title for s in result.sections]
    assert titles == ["Warranty Terms", "Coverage", "Claims"]
    assert "| SKU-4471 |" in result.sections[1].markdown.replace("  ", " ")
    assert result.sections[2].location["section"] == "Claims"
    assert result.raw and result.raw.get("schema_name") == "DoclingDocument"


def test_xlsx_sheets_become_sections(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "budget.xlsx"
    wb = Workbook()
    wb.active.title = "Cost centres"
    wb.active.append(["Code", "Owner"])
    wb.active.append(["CC-210", "Tom"])
    wb.create_sheet("Limits").append(["Role", "USD"])
    wb.save(path)

    result = extract_with_docling(path)
    assert [s.location.get("sheet") for s in result.sections] == [
        "Cost centres",
        "Limits",
    ]
    assert "pages" not in result.sections[0].location
    assert "CC-210" in result.sections[0].markdown


@pytest.mark.parametrize("status", ["partial_success", "failure"])
def test_incomplete_conversion_is_rejected(tmp_path, monkeypatch, status):
    from docling.datamodel.base_models import ConversionStatus

    from okf_mcp_server.src.ingest import docling_adapter

    result = SimpleNamespace(
        status=ConversionStatus(status),
        document=SimpleNamespace(
            iterate_items=lambda **kwargs: iter(()), export_to_dict=lambda: {}
        ),
    )
    monkeypatch.setattr(
        docling_adapter,
        "_converter",
        lambda _: SimpleNamespace(convert=lambda path: result),
    )
    with pytest.raises(ValueError, match="conversion did not complete"):
        extract_with_docling(tmp_path / "broken.pdf")
