"""End-to-end check of the Kubernetes or OpenShift deployment (`deployment/<overlay>`).

In a throwaway namespace it deploys Keycloak and PostgreSQL (test-only), then
the real manifests with generated secrets, and checks:

- pods run as an OpenShift-style arbitrary UID with a read-only root filesystem
- a Job created from the refresh CronJob syncs a mounted share and publishes
  a snapshot to the PVC; the server serves it read-only
- authentication and per-group filtering work in the cluster
- a second refresh after a source change is served without a restart
- tool calls and ingestion are audited

Usage:
  python tests/e2e/k8s_e2e.py [--context orbstack] [--image okf-mcp-server:dev]
  python tests/e2e/k8s_e2e.py --platform openshift --context <ctx> --build [--namespace <project>]

Kubernetes: needs kubectl and a cluster that can use the local image.
OpenShift: needs kubectl and oc logged in (`oc login`); `--build` builds the server
and ingest images in the cluster from the tracked files (BuildConfigs), the pods run
under restricted-v2 with the project's UID range, and the Route is checked. Use
`--namespace` where you cannot create projects (Developer Sandbox); only the
objects this test creates are deleted from it.
"""

import argparse
import asyncio
import json
import secrets
import socket
import subprocess
import sys
import time
from pathlib import Path

import yaml

sys.path.insert(0, str(Path(__file__).resolve().parent))
from auth_e2e import (  # noqa: E402  # noqa: E402
    CHECKS,
    KEYCLOAK_IMAGE,
    POSTGRES_IMAGE,
    check,
    mcp,
    realm,
)

ROOT = Path(__file__).resolve().parents[2]
UID = 1000650000
SCLORG_POSTGRES_IMAGE = "quay.io/sclorg/postgresql-15-c9s"  # runs as any UID
SOURCES = {
    "finance-travel.md": (
        "finance/travel.md",
        "# Travel\n\nThe hotel limit is 150 GBP per night.\n",
    ),
    "hr-salaries.md": (
        "hr/salaries.md",
        "# Salaries\n\nThe hotel allowance for relocation is 500 GBP.\n",
    ),
    "public-handbook.md": (
        "public/handbook.md",
        "# Handbook\n\nBook every hotel through the travel portal.\n",
    ),
}
RULES = "access:\n  share/finance/: [group:finance]\n  share/hr/: [group:hr]\n  share/public/: ['*']\n"


class Kube:
    def __init__(self, context: str, namespace: str):
        self.base = ["kubectl", "--context", context, "-n", namespace]

    def run(
        self, *args: str, stdin: str = "", check_rc: bool = True, timeout: int = 600
    ) -> str:
        out = subprocess.run(
            [*self.base, *args],
            input=stdin,
            capture_output=True,
            text=True,
            timeout=timeout,
        )
        if check_rc and out.returncode:
            raise RuntimeError(
                f"kubectl {' '.join(args[:3])}: {out.stderr.strip()[-800:]}"
            )
        return out.stdout.strip()

    def apply(self, docs: list) -> None:
        self.run("apply", "-f", "-", stdin=yaml.safe_dump_all(docs))

    def exec_py(self, code: str) -> str:
        return self.run(
            "exec", "deploy/okf-mcp-server", "--", "/app/.venv/bin/python", "-c", code
        )


def infra(
    passwords: dict,
    client_secret: str,
    admin_password: str,
    pg_password: str,
    openshift: bool = False,
) -> list:
    """Keycloak and PostgreSQL for the test (not part of the product manifests)."""
    realm_json = json.dumps(realm(client_secret, passwords))
    labels = lambda app: {"app": app}  # noqa: E731

    def deployment(
        name,
        image,
        env,
        ports,
        args=None,
        volumes=None,
        mounts=None,
        probe_path=None,
        probe_port=None,
    ):
        container = {
            "name": name,
            "image": image,
            "imagePullPolicy": "IfNotPresent",
            "env": [{"name": k, "value": v} for k, v in env.items()],
            "ports": [{"containerPort": p} for p in ports],
            "resources": {
                "requests": {"memory": "256Mi", "cpu": "100m"},
                "limits": {"memory": "1536Mi", "cpu": "1"},
            },
        }
        if args:
            container["args"] = args
        if mounts:
            container["volumeMounts"] = mounts
        if probe_path:
            container["readinessProbe"] = {
                "httpGet": {"path": probe_path, "port": probe_port},
                "periodSeconds": 3,
                "failureThreshold": 100,
            }
        return {
            "apiVersion": "apps/v1",
            "kind": "Deployment",
            "metadata": {"name": name, "labels": labels(name)},
            "spec": {
                "selector": {"matchLabels": labels(name)},
                "template": {
                    "metadata": {"labels": labels(name)},
                    "spec": {"containers": [container], "volumes": volumes or []},
                },
            },
        }

    def service(name, port):
        return {
            "apiVersion": "v1",
            "kind": "Service",
            "metadata": {"name": name},
            "spec": {
                "selector": labels(name.replace("postgresql", "postgres")),
                "ports": [{"port": port, "targetPort": port}],
            },
        }

    return [
        {
            "apiVersion": "v1",
            "kind": "Secret",
            "metadata": {"name": "keycloak-realm"},
            "stringData": {"okf.json": realm_json},
        },
        deployment(
            "keycloak",
            KEYCLOAK_IMAGE,
            {
                "KC_BOOTSTRAP_ADMIN_USERNAME": "admin",
                "KC_BOOTSTRAP_ADMIN_PASSWORD": admin_password,
                "KC_HEALTH_ENABLED": "true",
            },
            [8080, 9000],
            args=["start-dev", "--import-realm"],
            volumes=[{"name": "realm", "secret": {"secretName": "keycloak-realm"}}],
            mounts=[
                {
                    "name": "realm",
                    "mountPath": "/opt/keycloak/data/import",
                    "readOnly": True,
                }
            ],
            probe_path="/health/ready",
            probe_port=9000,
        ),
        service("keycloak", 8080),
        deployment(
            "postgres",
            SCLORG_POSTGRES_IMAGE if openshift else POSTGRES_IMAGE,
            {
                "POSTGRESQL_USER": "okf",
                "POSTGRESQL_PASSWORD": pg_password,
                "POSTGRESQL_DATABASE": "okf",
            }
            if openshift
            else {
                "POSTGRES_USER": "okf",
                "POSTGRES_PASSWORD": pg_password,
                "POSTGRES_DB": "okf",
            },
            [5432],
        ),
        service("postgresql", 5432),
    ]


def docx_bytes() -> bytes:
    """A small Word document for the ingest image to convert with Docling."""
    import io

    from docx import Document

    doc = Document()
    doc.add_heading("Finance Hotel Guide", 0)
    doc.add_heading("London", 1)
    doc.add_paragraph("In London the hotel limit is 260 GBP per night.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def product(
    image: str,
    client_secret: str,
    pg_password: str,
    sources: dict,
    ingest_image: str = "",
    platform: str = "kubernetes",
    docx: bool = False,
) -> list:
    """The deployment/<platform> overlay, filled in for this cluster.

    An empty `image` keeps the overlay's images (OpenShift: the ImageStream tags
    built in the cluster). On OpenShift the UID is left to the SCC.
    """
    docx = docx or bool(ingest_image)
    rendered = subprocess.run(
        ["kubectl", "kustomize", str(ROOT / "deployment" / platform)],
        capture_output=True,
        text=True,
        check=True,
    ).stdout
    docs = [d for d in yaml.safe_load_all(rendered) if d]
    kc = "http://keycloak:8080/realms/okf/protocol/openid-connect"
    for doc in docs:
        kind, meta = doc["kind"], doc["metadata"]["name"]
        if kind == "ConfigMap" and meta == "okf-mcp-server-config":
            doc["data"].update(
                {
                    "SSO_INTROSPECTION_URL": f"{kc}/token/introspect",
                    "SSO_AUTHORIZATION_URL": f"{kc}/auth",
                    "SSO_TOKEN_URL": f"{kc}/token",
                    "SSO_CALLBACK_URL": "http://localhost/auth/callback/oidc",
                    "MCP_HOST_ENDPOINT": "http://localhost",
                    "ENVIRONMENT": "development",
                }
            )
        elif kind == "Secret" and meta == "okf-mcp-server-secrets":
            doc["stringData"].update(
                {
                    "SSO_CLIENT_SECRET": client_secret,
                    "SESSION_SECRET": secrets.token_urlsafe(32),
                    "POSTGRES_USER": "okf",
                    "POSTGRES_PASSWORD": pg_password,
                }
            )
        elif kind == "ConfigMap" and meta == "okf-connectors":
            cfg = yaml.safe_load(doc["data"]["connectors.yaml"])
            cfg["okf_config"] = "/etc/okf/_okf.yaml"
            doc["data"] = {"connectors.yaml": yaml.safe_dump(cfg), "_okf.yaml": RULES}
        if kind in ("Deployment", "CronJob"):
            spec = (
                doc["spec"]["template"]["spec"]
                if kind == "Deployment"
                else doc["spec"]["jobTemplate"]["spec"]["template"]["spec"]
            )
            if platform == "kubernetes":
                spec["securityContext"].update({"runAsUser": UID, "runAsGroup": 0})
            for c in spec["containers"]:
                if kind == "CronJob" and ingest_image:
                    c["image"] = ingest_image
                elif image:
                    c["image"] = image
            if kind == "CronJob":
                for v in spec["volumes"]:
                    if v["name"] == "sources":
                        items = [{"key": k, "path": p} for k, (p, _) in sources.items()]
                        if docx:
                            items.append(
                                {
                                    "key": "finance-hotel-guide.docx",
                                    "path": "finance/hotel-guide.docx",
                                }
                            )
                        v.clear()
                        v.update(
                            {
                                "name": "sources",
                                "configMap": {
                                    "name": "okf-e2e-sources",
                                    "items": items,
                                },
                            }
                        )
    source_map = {
        "apiVersion": "v1",
        "kind": "ConfigMap",
        "metadata": {"name": "okf-e2e-sources"},
        "data": {k: text for k, (_, text) in sources.items()},
    }
    if docx:
        import base64

        source_map["binaryData"] = {
            "finance-hotel-guide.docx": base64.b64encode(docx_bytes()).decode()
        }
    docs.append(source_map)
    return docs


def source_archive() -> Path:
    """Tracked files incl. uncommitted changes (never .venv or data/), for oc start-build."""
    ref = subprocess.run(
        ["git", "stash", "create"], cwd=ROOT, capture_output=True, text=True
    ).stdout.strip()
    archive = Path("/tmp") / f"okf-e2e-src-{secrets.token_hex(3)}.tar.gz"
    subprocess.run(
        ["git", "archive", "--format=tar.gz", "-o", str(archive), ref or "HEAD"],
        cwd=ROOT,
        check=True,
    )
    return archive


def free_port() -> int:
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--context", default="orbstack")
    parser.add_argument(
        "--platform", choices=("kubernetes", "openshift"), default="kubernetes"
    )
    parser.add_argument(
        "--image",
        default=None,
        help="server image (default: okf-mcp-server:dev; OpenShift: the ImageStream)",
    )
    parser.add_argument(
        "--namespace", help="use this existing namespace/project instead of a new one"
    )
    parser.add_argument(
        "--build",
        action="store_true",
        help="OpenShift: build the server and ingest images with the BuildConfigs",
    )
    parser.add_argument(
        "--keep", action="store_true", help="keep the namespace for inspection"
    )
    parser.add_argument(
        "--ingest-image",
        default="",
        help="run the refresh CronJob with this image and add a DOCX source",
    )
    args = parser.parse_args()
    openshift = args.platform == "openshift"
    if args.image is None:
        args.image = "" if openshift else "okf-mcp-server:dev"
    if openshift and not (args.build or args.image):
        parser.error("--platform openshift needs --build or --image")
    docx = bool(args.ingest_image) or (openshift and args.build)
    namespace = args.namespace or f"okf-e2e-{secrets.token_hex(3)}"
    passwords = {u: secrets.token_urlsafe(12) for u in ("alice", "bob", "carol")}
    client_secret, admin_password, pg_password = (
        secrets.token_urlsafe(18) for _ in range(3)
    )
    kube = Kube(args.context, namespace)
    forward = None
    applied: list = []

    def render(sources: dict) -> list:
        return product(
            args.image,
            client_secret,
            pg_password,
            sources,
            args.ingest_image,
            args.platform,
            docx,
        )

    try:
        if not args.namespace:
            create = (
                [
                    "oc",
                    "--context",
                    args.context,
                    "new-project",
                    namespace,
                    "--skip-config-write",
                ]
                if openshift
                else [
                    "kubectl",
                    "--context",
                    args.context,
                    "create",
                    "namespace",
                    namespace,
                ]
            )
            subprocess.run(create, check=True, capture_output=True)
        print(f"namespace {namespace}: deploying Keycloak and PostgreSQL ...")
        applied += infra(
            passwords, client_secret, admin_password, pg_password, openshift
        )
        kube.apply(applied)
        kube.run("rollout", "status", "deploy/keycloak", "--timeout=420s", timeout=460)
        kube.run("rollout", "status", "deploy/postgres", "--timeout=180s", timeout=200)

        docs = render(SOURCES)
        if openshift and args.build:
            builds = [d for d in docs if d["kind"] in ("ImageStream", "BuildConfig")]
            applied += builds
            kube.apply(builds)
            archive = source_archive()
            try:
                for bc in ("okf-mcp-server", "okf-mcp-server-ingest"):
                    print(
                        f"building {bc} in the cluster (the ingest image takes a while) ..."
                    )
                    out = subprocess.run(
                        [
                            "oc",
                            "--context",
                            args.context,
                            "-n",
                            namespace,
                            "start-build",
                            bc,
                            f"--from-archive={archive}",
                            "--wait",
                        ],
                        capture_output=True,
                        text=True,
                        timeout=3600,
                    )
                    check(
                        f"in-cluster build {bc}", out.returncode == 0, out.stderr[-800:]
                    )
                    if out.returncode:
                        return 1
            finally:
                archive.unlink(missing_ok=True)

        print("deploying the okf manifests ...")
        applied += docs
        kube.apply(docs)
        kube.run("create", "job", "--from=cronjob/okf-refresh", "refresh-1")
        done = kube.run(
            "wait",
            "--for=condition=complete",
            "job/refresh-1",
            "--timeout=300s",
            check_rc=False,
            timeout=320,
        )
        logs = kube.run("logs", "job/refresh-1", check_rc=False)
        check(
            "refresh Job (from the CronJob) syncs the share and publishes",
            "condition met" in done
            and f"synced share: {4 if docx else 3} downloaded" in logs
            and "published" in logs
            and "FAILED" not in logs,
            logs[-600:],
        )
        kube.run(
            "rollout", "status", "deploy/okf-mcp-server", "--timeout=300s", timeout=320
        )

        uid = kube.run("exec", "deploy/okf-mcp-server", "--", "id", "-u")
        if openshift:
            # restricted-v2 assigns a UID from the project's range, e.g. "1000650000/10000".
            uid_range = kube.run(
                "get",
                "namespace",
                namespace,
                "-o",
                "jsonpath={.metadata.annotations.openshift\\.io/sa\\.scc\\.uid-range}",
                check_rc=False,
            )
            start, _, size = uid_range.partition("/")
            ok = (
                int(start) <= int(uid) < int(start) + int(size)
                if start.isdigit() and size.isdigit()
                else int(uid) >= 1000000000
            )
            check(
                f"server runs with a UID from the SCC range ({uid_range or '?'})",
                ok,
                uid,
            )
            host = kube.run(
                "get", "route/okf-mcp-server", "-o", "jsonpath={.spec.host}"
            )
            import httpx

            # verify=False: clusters often use a self-signed router certificate.
            health = httpx.get(f"https://{host}/health", verify=False, timeout=15)
            unauth = httpx.post(
                f"https://{host}/mcp",
                json={"jsonrpc": "2.0", "id": 1, "method": "ping"},
                headers={"Accept": "application/json, text/event-stream"},
                verify=False,
                timeout=15,
            )
            check(
                f"Route https://{host} serves /health and requires a token on /mcp",
                health.status_code == 200 and unauth.status_code == 401,
                f"{health.status_code} {unauth.status_code}",
            )
        else:
            check(f"server runs as arbitrary UID {UID}", uid == str(UID), uid)
        ro = kube.run(
            "exec",
            "deploy/okf-mcp-server",
            "--",
            "sh",
            "-c",
            "awk '$2==\"/\"{print $4}' /proc/mounts",
            check_rc=False,
        )
        check("root filesystem is read-only", "ro" in ro.split(","), ro)
        cron = json.loads(kube.run("get", "cronjob/okf-refresh", "-o", "json"))
        check(
            "refresh CronJob is scheduled, one run at a time",
            cron["spec"]["schedule"] == "0 2 * * *"
            and cron["spec"]["concurrencyPolicy"] == "Forbid",
        )

        def token(user: str) -> str:
            return kube.exec_py(
                "import httpx;print(httpx.post('http://keycloak:8080/realms/okf/protocol/openid-connect/token',"
                f"data={{'grant_type':'password','client_id':'okf-cli','username':'{user}',"
                f"'password':'{passwords[user]}','scope':'openid'}}).json()['access_token'])"
            )

        port = free_port()
        forward = subprocess.Popen(
            [*kube.base, "port-forward", "svc/okf-mcp-server", f"{port}:5001"],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
        url = f"http://localhost:{port}/mcp"
        import httpx

        for _ in range(60):
            try:
                if (
                    httpx.get(f"http://localhost:{port}/health", timeout=2).status_code
                    == 200
                ):
                    break
            except httpx.HTTPError:
                time.sleep(1)
        init = {
            "jsonrpc": "2.0",
            "id": 1,
            "method": "initialize",
            "params": {
                "protocolVersion": "2025-06-18",
                "capabilities": {},
                "clientInfo": {"name": "k8s-e2e", "version": "1"},
            },
        }
        r = httpx.post(
            url,
            json=init,
            headers={"Accept": "application/json, text/event-stream"},
            timeout=10,
        )
        check(
            "no token -> 401 in the cluster", r.status_code == 401, str(r.status_code)
        )

        def visible(bearer: str) -> set:
            result = asyncio.run(
                mcp(url, bearer, "search_knowledge", {"query": "hotel", "limit": 20})
            )
            return {x["concept_id"] for x in result.structured_content["results"]}

        alice, bob = token("alice"), token("bob")
        seen = visible(alice)
        finance = {"share/finance/travel", "share/public/handbook"}
        if docx:
            finance.add(
                "share/finance/hotel-guide"
            )  # the DOCX, converted by Docling in the Job
        check(
            "alice (finance) sees finance + public"
            + (" incl. the Docling-converted DOCX" if docx else ""),
            seen == finance,
            str(seen),
        )
        seen = visible(bob)
        check(
            "bob (hr) sees hr + public",
            seen == {"share/hr/salaries", "share/public/handbook"},
            str(seen),
        )

        # Change the source and run the scheduled refresh again: no server restart.
        sources = dict(SOURCES)
        sources["hr-leave.md"] = (
            "hr/leave.md",
            "# Leave\n\nEvery hotel stay during leave is unpaid.\n",
        )
        kube.apply(
            [
                d
                for d in render(sources)
                if d["kind"] in ("CronJob", "ConfigMap")
                and d["metadata"]["name"] in ("okf-refresh", "okf-e2e-sources")
            ]
        )
        kube.run("create", "job", "--from=cronjob/okf-refresh", "refresh-2")
        kube.run(
            "wait",
            "--for=condition=complete",
            "job/refresh-2",
            "--timeout=300s",
            check_rc=False,
            timeout=320,
        )
        seen = visible(token("bob"))
        check(
            "second refresh is served without a restart",
            "share/hr/leave" in seen,
            str(seen),
        )

        server_logs = kube.run("logs", "deploy/okf-mcp-server", check_rc=False)
        check(
            "tool calls are audited with the caller",
            "okf.audit" in server_logs and "alice@example.com" in server_logs,
        )
        ingest_audit = kube.run(
            "exec",
            "deploy/okf-mcp-server",
            "--",
            "cat",
            "/app/data/ingest-audit.jsonl",
            check_rc=False,
        )
        events = [
            json.loads(line)["event"]
            for line in ingest_audit.splitlines()
            if line.strip()
        ]
        check(
            "ingestion is audited on the volume",
            events.count("ingest.build") == 2 and "ingest.sync" in events,
            str(events),
        )
    finally:
        if forward:
            forward.terminate()
        if args.keep:
            print(f"kept namespace {namespace}")
        elif args.namespace:  # shared project: delete only what this test created
            kube.run(
                "delete",
                "-f",
                "-",
                "--ignore-not-found",
                "--wait=false",
                stdin=yaml.safe_dump_all(applied),
                check_rc=False,
            )
            kube.run(
                "delete",
                "job,build" if openshift else "job",
                "-l",
                "app=okf-mcp-server",
                "--ignore-not-found",
                "--wait=false",
                check_rc=False,
            )
        else:
            subprocess.run(
                [
                    "kubectl",
                    "--context",
                    args.context,
                    "delete",
                    "namespace",
                    namespace,
                    "--wait=false",
                ],
                capture_output=True,
            )
    failed = [name for name, ok in CHECKS if not ok]
    print(f"\n{len(CHECKS) - len(failed)}/{len(CHECKS)} checks passed")
    return 1 if failed or not CHECKS else 0


if __name__ == "__main__":
    sys.exit(main())
